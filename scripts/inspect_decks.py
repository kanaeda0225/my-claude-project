"""Programmatic visual-design inspection of generated .pptx decks.

Since we can't render to images in this sandbox, inspect the slide
structure programmatically and report:
- Each slide's title, subtitle, kind, content count
- Detected layout violations (text too long, fonts too small etc.)
- Quick "table of contents" summary
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt


def inspect_deck(path: Path):
    print(f"\n{'='*70}")
    print(f"FILE: {path.name}")
    print('='*70)
    pres = Presentation(path)
    print(f"Slides: {len(pres.slides)}")
    print(f"Size: {pres.slide_width/914400:.1f} x {pres.slide_height/914400:.1f} inches")
    print()

    for i, slide in enumerate(pres.slides):
        title_text = ""
        subtitle_text = ""
        body_chars = 0
        n_textbox = 0
        smallest_font = None
        has_table = False
        has_ascii = False

        for shape in slide.shapes:
            if shape.has_text_frame:
                tf_text = shape.text_frame.text
                if not tf_text.strip():
                    continue
                n_textbox += 1
                body_chars += len(tf_text)
                # Try to detect smallest font in this textbox
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        if run.font.size:
                            sz = run.font.size.pt
                            if smallest_font is None or sz < smallest_font:
                                smallest_font = sz
                # Try heuristics for title/subtitle
                first_para = shape.text_frame.paragraphs[0]
                if first_para.runs:
                    first_size = first_para.runs[0].font.size
                    if first_size and first_size.pt >= 28 and not title_text:
                        title_text = tf_text.strip().split("\n")[0]
                    elif first_size and 18 <= first_size.pt <= 24 and not subtitle_text:
                        subtitle_text = tf_text.strip().split("\n")[0]
                # ASCII detection
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        if run.font.name == "Courier New":
                            has_ascii = True
            if shape.shape_type == 19:  # MSO_SHAPE_TYPE.TABLE
                has_table = True

        marker = ""
        if i == 0:
            marker = "[COVER]"
        if has_table:
            marker += " [TABLE]"
        if has_ascii:
            marker += " [ASCII]"

        # Detect potential issues
        warnings = []
        if smallest_font is not None and smallest_font < 13:
            warnings.append(f"⚠️ smallest font {smallest_font}pt (<13)")
        if body_chars > 600:
            warnings.append(f"⚠️ very text-heavy ({body_chars} chars)")

        title_disp = title_text[:50] if title_text else "(no title detected)"
        print(f"  {i+1:3}: {title_disp:50} {marker}")
        if subtitle_text:
            print(f"      └ {subtitle_text[:80]}")
        for w in warnings:
            print(f"      {w}")


if __name__ == "__main__":
    paths = [
        Path("outputs/4971-mec-slides.pptx"),
        Path("outputs/421A-movin-slides.pptx"),
        Path("outputs/6648-kawaden-slides.pptx"),
        Path("outputs/5537-albalink-slides.pptx"),
        Path("outputs/4112-hodogaya-slides.pptx"),
    ]
    for p in paths:
        if p.exists():
            inspect_deck(p)
